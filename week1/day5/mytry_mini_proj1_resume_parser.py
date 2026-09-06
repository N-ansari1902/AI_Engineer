import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("API key not initialized")

client = Groq(api_key = my_api_key)

import pdfplumber
from docx import Document


def extract_from_pdf(file_path):
    all_text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                all_text += text + "\n"
            for table in page.extract_tables():
                for row in table:
                    row_text = " | ".join(cell.strip() for cell in row if cell and cell.strip())
                    if row_text:
                        all_text += row_text + "\n"
    return all_text.strip()


def extract_from_docx(file_path):
    doc = Document(file_path)
    all_text = ""
    for para in doc.paragraphs:
        if para.text.strip():
            all_text += para.text.strip() + "\n"
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                all_text += row_text + "\n"
    for section in doc.sections:
        for para in section.header.paragraphs:
            if para.text.strip():
                all_text += para.text.strip() + "\n"
        for para in section.footer.paragraphs:
            if para.text.strip():
                all_text += para.text.strip() + "\n"
    return all_text.strip()


def extract_resume_text(filename):
    # builds the full path by joining the folder this script is in + the filename
    folder = os.path.dirname(os.path.abspath(__file__))
    file_path = os.path.join(folder, filename)

    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return extract_from_pdf(file_path)
    elif ext == ".docx":
        return extract_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported format: '{ext}'. Use .pdf or .docx only.")


# ── Usage ──
resume_text = extract_resume_text("resume.pdf")   # or "resume.docx"
job_description = extract_resume_text("job_jd.pdf")

model = "openai/gpt-oss-120b"
role = "user"
from pydantic import BaseModel
class Candidate(BaseModel):
    skills:str
    certifications:str

schema=Candidate.model_json_schema()

# Add this after your Candidate class
class MatchResult(BaseModel):
    percentage_match: int
    match_summary: str

match_schema = MatchResult.model_json_schema()

response_format = {
    "type" : "json_object"
}

system_prompt = f"""
Extract the information of candidates strictly based on this schema format into json.
{schema}
"""

message_system = {
    "role" : "system",
    "content" : system_prompt
}

prompt = f"""
This is a candidates resume. Please extract all the information from this.
{resume_text}
"""
message = {
    "role" : role,
    "content" : prompt
}


messages = [message_system, message]

response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)

answer = response.choices[0].message.content

import json
raw_json = answer
data_file = json.loads(raw_json)
candidate_details=Candidate(**data_file)


matching_prompt = f"""
Match the candidate's skills {candidate_details.skills} with the job description given {job_description} and tell the percentage match of the candidate's profile for that job, strictly based on this schema format into json.
{match_schema}
"""

matching_system = {
    "role" : "system",
    "content" : matching_prompt
}

final_mes = [
    matching_system,
    {"role": "user", 
     "content": "Provide the match result based on the above instructions."
    }
]

final_act = client.chat.completions.create(model=model, messages=final_mes, response_format=response_format)
final_res = final_act.choices[0].message.content

match_data = json.loads(final_res)
match_result = MatchResult(**match_data)

if match_result.percentage_match >= 70:
    print("Shortlisted!")
    print(f"Match: {match_result.percentage_match}%")
    print(f"Summary: {match_result.match_summary}")

else:
    print("Rejected! Minimum matching criterion not fulfilled.")
    print(f"Summary: {match_result.match_summary}")